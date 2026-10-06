"""修订执行：在副本上改，原文件不动，改前备份，改后回归。

铁律（用户明确要求）：
  1. 绝不修改原文件
  2. 改动一律输出到新文件
  3. 自动备份

实现要点（来自 ~/.workbuddy/skills/docx-section-revision/SKILL.md）：
  - 先把 Finding 解析为带 expected_old_text 的 EditProposal，再做冲突/源 hash preflight
  - 只改动命中的 w:t 节点，保留超链接、公式、域、书签和批注锚点等结构化 XML
  - 输出到新文件并在发布前完成 DOCX 保存和确定性回归复验
"""

from __future__ import annotations

import json
import hashlib
import difflib
import shutil
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable

import docx
from docx.text.paragraph import Paragraph

from paperaudit.evidence import run_all
from paperaudit.ingest import read_docx
from paperaudit.reporting import apply_gate

_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_W_T = f"{_NS}t"
_W_DEL_TEXT = f"{_NS}delText"


def _index_paragraphs(document) -> dict[str, Paragraph]:
    """按 body 遍历序重建 block_id → 段落 的映射。

    必须与 ingest 的编号规则一致（p_{序号:04d}，只数 w:p）。
    """
    idx: dict[str, Paragraph] = {}
    n = 0
    for el in document.element.body.iterchildren():
        if el.tag.split("}")[-1] == "p":
            n += 1
            idx[f"p_{n:04d}"] = Paragraph(el, document)
    return idx


def apply_revision(
    source: str | Path,
    run_dir: str | Path,
    finding_ids: list[str],
    out_path: str | Path | None = None,
    backup: bool = True,
    text: str | None = None,
    progress: Callable[..., None] | None = None,
) -> dict:
    src = Path(source)
    run_dir = Path(run_dir)
    if not src.exists():
        raise FileNotFoundError(src)
    if src.suffix.casefold() != ".docx":
        return {
            "status": "error",
            "message": "安全修订目前仅支持 DOCX；PDF 可直接审查，但必须先在可编辑 DOCX 上应用修改。",
        }

    _progress(progress, "load", "正在加载已接受的修改意见", 10)

    findings = _load_findings(run_dir)
    fixable = [f.get("id") for f in findings if _proposal_text(f)]

    if finding_ids == ["*"]:
        finding_ids = fixable  # 一键应用所有带改写内容的条目

    selected = [f for f in findings if f.get("id") in finding_ids]
    if not selected:
        return {
            "status": "error",
            "message": f"未找到指定条目：{finding_ids}",
            "fixable": fixable,
            "hint": "用 --findings '*' 应用所有带改写内容的条目，或用 --text 直接给出文本",
        }

    src = src.resolve()
    source_digest = hashlib.sha256(src.read_bytes()).hexdigest()
    manifest_path = run_dir / "manifest.json"
    manifest: dict[str, Any] = {}
    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            manifest = {}
        prepared_hash = str(manifest.get("hash") or "") if isinstance(manifest, dict) else ""
        if prepared_hash and prepared_hash != source_digest[:16]:
            return {
                "status": "error",
                "message": "源文件与审查包的 hash 不一致，请重新 prepare 和 verify 后再修改。",
                "applied": [],
                "failed": [{"id": None, "reason": "source-hash-mismatch"}],
                "output": None,
                "backup": None,
            }

    out = Path(out_path) if out_path else src.with_name(f"{src.stem}_revised{src.suffix}")
    out = out.resolve()
    if out == src:
        return {
            "status": "error",
            "message": "修订输出路径不能覆盖源文件。",
            "applied": [],
            "failed": [{"id": None, "reason": "output-is-source"}],
            "output": None,
            "backup": None,
        }
    if out.exists():
        return {
            "status": "error",
            "message": f"输出文件已存在，为避免覆盖而停止：{out}",
            "applied": [],
            "failed": [{"id": None, "reason": "output-already-exists"}],
            "output": str(out),
            "backup": None,
        }

    # Plan and validate every selected edit against the original source before
    # creating a backup or touching the output path.  Conflicting proposals
    # are rejected as a group, independent of their order in findings.json.
    document = docx.Document(str(src))
    paras = _index_paragraphs(document)
    planned, failed = _preflight_proposals(selected, paras, text)
    revision_plan = _revision_plan_payload(
        src,
        run_dir,
        finding_ids,
        selected,
        planned,
        failed,
        source_digest,
        str(manifest.get("hash") or ""),
        fixable,
        paras,
    )
    revision_plan["revision_plan"] = str(run_dir / "revision.plan.json")
    _persist_revision_plan(run_dir, revision_plan)
    if not planned:
        # A blocked plan must not create a file that looks like a successful
        # revision.  The plan and diff above are the review artifacts; callers
        # can retry after resolving the failed proposals.
        return {
            "status": "error",
            "message": "没有可安全应用的修订建议。",
            "applied": [],
            "failed": failed,
            "output": None,
            "backup": None,
            "fixable": fixable,
            "revision_plan": str(run_dir / "revision.plan.json"),
            "diff": str(run_dir / "revision.diff"),
        }

    # Detect concurrent source changes since parsing, then preserve an
    # immutable backup with a collision-resistant name.
    if hashlib.sha256(src.read_bytes()).hexdigest() != source_digest:
        return {
            "status": "error",
            "message": "修订期间源文件发生变化，请重新 prepare 后再试。",
            "applied": [],
            "failed": [{"id": None, "reason": "source-changed-during-preflight"}],
            "output": None,
            "backup": None,
        }
    backup_path = None
    if backup:
        ts = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        backup_path = src.with_suffix(f".bak-{ts}{src.suffix}")
        shutil.copy2(src, backup_path)
    _progress(progress, "backup", "正在创建原文备份和修改副本", 30)

    applied = []
    for f, operations in planned:
        for bid, start, end, replacement in operations:
            _replace_span_keep_xml(paras[bid], start, end, replacement)
            applied.append({"id": f.get("id"), "block": bid})

    out.parent.mkdir(parents=True, exist_ok=True)
    temporary_out = out.with_name(f".{out.stem}.paperaudit-{uuid.uuid4().hex}.tmp.docx")
    try:
        document.save(str(temporary_out))
        # Publish only a complete DOCX package.  A failed save leaves the
        # caller's intended output path untouched.
        temporary_out.replace(out)
    finally:
        if temporary_out.exists():
            temporary_out.unlink()

    # 回归复验：改完重跑确定性检查，确认没引入新问题
    before_findings = findings
    after_ir = read_docx(out)
    after = apply_gate(after_ir, run_all(after_ir))
    regression_report = _compare_findings(before_findings, [f.to_dict() for f in after if f.gate_passed])
    regressions = regression_report["regressions"]
    _progress(
        progress,
        "regression",
        "正在执行修改后的回归复验",
        92,
        regressions=len(regressions),
        resolved=len(regression_report["resolved"]),
        unchanged=len(regression_report["unchanged"]),
    )

    result = {
        "status": "ok" if applied else "error",
        "output": str(out),
        "backup": str(backup_path) if backup_path else None,
        "applied": applied,
        "failed": failed,
        "regressions": regressions,
        "regression_report": regression_report,
        "source_unchanged": str(src),
        "fixable": fixable,
        "revision_plan": str(run_dir / "revision.plan.json"),
        "diff": str(run_dir / "revision.diff"),
    }
    if not applied:
        result["hint"] = "所选条目均无改写内容；可用 --text 直接给出文本，或 --findings '*'"
    return result


def build_revision_plan(
    source: str | Path,
    run_dir: str | Path,
    finding_ids: list[str],
    *,
    text: str | None = None,
    persist: bool = True,
) -> dict[str, Any]:
    """Create an explicit, reviewable edit plan without changing the DOCX.

    Every ready proposal contains target block spans, the exact old text and
    replacement, plus the optimistic-lock precondition used by
    :func:`apply_revision`. Conflicts and unsafe level-C proposals remain in
    ``failed``/``conflict_graph`` so a UI can show why approval is blocked.
    """

    src = Path(source).resolve()
    run = Path(run_dir).resolve()
    if not src.exists():
        raise FileNotFoundError(src)
    if src.suffix.casefold() != ".docx":
        plan = {
            "schema_version": 1,
            "status": "unsupported",
            "message": "安全修订目前仅支持 DOCX；PDF 可审查但不能生成 DOCX 修订计划。",
            "source": str(src),
            "run_dir": str(run),
            "proposals": [],
            "failed": [],
        }
        if persist:
            plan["revision_plan"] = str(run / "revision.plan.json")
            _persist_revision_plan(run, plan)
        return plan

    findings = _load_findings(run)
    fixable = [f.get("id") for f in findings if isinstance(f, dict) and _proposal_text(f)]
    requested = [str(value) for value in finding_ids]
    if requested == ["*"]:
        requested = [str(value) for value in fixable]
    requested_set = set(requested)
    selected = [f for f in findings if isinstance(f, dict) and str(f.get("id")) in requested_set]
    source_digest = hashlib.sha256(src.read_bytes()).hexdigest()
    manifest = _read_json_object(run / "manifest.json")
    prepared_hash = str(manifest.get("hash") or "")
    if not selected:
        plan = {
            "schema_version": 1,
            "status": "blocked",
            "message": f"未找到指定条目：{requested}",
            "source": str(src),
            "run_dir": str(run),
            "source_hash": source_digest,
            "prepared_hash": prepared_hash,
            "source_hash_matches": not prepared_hash or prepared_hash == source_digest[:16],
            "requested_finding_ids": requested,
            "fixable": fixable,
            "proposals": [],
            "failed": [{"id": requested, "reason": "finding-not-found"}],
            "conflict_graph": {"nodes": [], "edges": []},
            "generated_at": datetime.now(timezone.utc).isoformat(),
        }
        if persist:
            plan["revision_plan"] = str(run / "revision.plan.json")
            _persist_revision_plan(run, plan)
        return plan

    document = docx.Document(str(src))
    paras = _index_paragraphs(document)
    planned, failed = _preflight_proposals(selected, paras, text)
    plan = _revision_plan_payload(
        src,
        run,
        requested,
        selected,
        planned,
        failed,
        source_digest,
        prepared_hash,
        fixable,
        paras,
    )
    if prepared_hash and prepared_hash != source_digest[:16]:
        plan["status"] = "blocked"
        plan["failed"].append({"id": None, "reason": "source-hash-mismatch"})
        plan["message"] = "源文件与审查包 hash 不一致，请重新 prepare 和 verify。"
    if persist:
        plan["revision_plan"] = str(run / "revision.plan.json")
        _persist_revision_plan(run, plan)
    return plan


def _revision_plan_payload(
    source: Path,
    run: Path,
    requested: list[str],
    selected: list[dict],
    planned: list[tuple[dict, list[tuple[str, int, int, str]]]],
    failed: list[dict[str, Any]],
    source_digest: str,
    prepared_hash: str,
    fixable: list[Any],
    paragraphs: dict[str, Paragraph],
) -> dict[str, Any]:
    proposals: list[dict[str, Any]] = []
    for finding, operations in planned:
        proposal = _proposal(finding)
        proposal_id = str(proposal.get("id") or f"edit-{finding.get('id', uuid.uuid4().hex[:8])}")
        level = str(proposal.get("level") or finding.get("edit_level") or "B_meaning_preserving")
        operation_rows = []
        for block_id, start, end, replacement in operations:
            actual = _paragraph_text(paragraphs[block_id])
            operation_rows.append(
                {
                    "block_id": block_id,
                    "start": start,
                    "end": end,
                    "expected_old_text": actual[start:end],
                    "replacement": replacement,
                }
            )
        proposals.append(
            {
                "id": proposal_id,
                "finding_ids": [str(finding.get("id"))],
                "target_blocks": [row["block_id"] for row in operation_rows],
                "edit_type": str(proposal.get("edit_type") or finding.get("issue_type") or "revision"),
                "level": level,
                "old_text": "\n".join(row["expected_old_text"] for row in operation_rows),
                "new_text": "\n".join(row["replacement"] for row in operation_rows),
                "rationale": str(proposal.get("rationale") or finding.get("rationale") or ""),
                "risk_flags": _risk_flags(proposal.get("risk_flags"), len(operation_rows)),
                "operations": operation_rows,
            }
        )

    nodes = [{"id": str(finding.get("id")), "status": "ready"} for finding, _ in planned]
    edges: list[dict[str, Any]] = []
    for item in failed:
        ids = item.get("id") if isinstance(item, dict) else None
        if not isinstance(ids, list) or len(ids) < 2 or "冲突" not in str(item.get("reason", "")):
            continue
        for index, left in enumerate(ids):
            for right in ids[index + 1 :]:
                edges.append({"from": str(left), "to": str(right), "type": "target_block_conflict", "reason": item.get("reason", "")})
    for item in failed:
        if not isinstance(item, dict):
            continue
        ids = item.get("id") if isinstance(item.get("id"), list) else [item.get("id")]
        for identifier in ids:
            if identifier is not None and not any(node.get("id") == str(identifier) for node in nodes):
                nodes.append({"id": str(identifier), "status": "blocked", "reason": item.get("reason", "")})
    status = "ready" if proposals and not failed else ("partial" if proposals else "blocked")
    diff_entries = _diff_entries(proposals)
    return {
        "schema_version": 1,
        "status": status,
        "source": str(source),
        "run_dir": str(run),
        "source_hash": source_digest,
        "prepared_hash": prepared_hash,
        "source_hash_matches": not prepared_hash or prepared_hash == source_digest[:16],
        "requested_finding_ids": requested,
        "selected_finding_ids": [str(f.get("id")) for f in selected],
        "fixable": [str(value) for value in fixable],
        "proposals": proposals,
        "diff": diff_entries,
        "diff_path": str(run / "revision.diff"),
        "failed": failed,
        "conflict_graph": {"nodes": nodes, "edges": edges},
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def _read_json_object(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return {}
    return value if isinstance(value, dict) else {}


def _risk_flags(value: Any, operation_count: int) -> list[str]:
    if isinstance(value, str):
        flags = [value] if value.strip() else []
    elif isinstance(value, (list, tuple, set)):
        flags = [str(item) for item in value if str(item).strip()]
    else:
        flags = []
    if operation_count > 1 and "multi_block" not in flags:
        flags.append("multi_block")
    return flags


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _diff_entries(proposals: list[dict[str, Any]]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for proposal in proposals:
        for operation in proposal.get("operations", []) or []:
            before = str(operation.get("expected_old_text", ""))
            after = str(operation.get("replacement", ""))
            unified = "\n".join(
                difflib.unified_diff(
                    before.splitlines(),
                    after.splitlines(),
                    fromfile=f"{operation.get('block_id', '')}:before",
                    tofile=f"{operation.get('block_id', '')}:after",
                    lineterm="",
                )
            )
            entries.append(
                {
                    "proposal_id": proposal.get("id", ""),
                    "finding_ids": proposal.get("finding_ids", []),
                    "block_id": operation.get("block_id", ""),
                    "before": before,
                    "after": after,
                    "unified": unified,
                }
            )
    return entries


def _persist_revision_plan(run: Path, plan: dict[str, Any]) -> None:
    """Persist both machine-readable proposal data and a text diff artifact."""

    diff_path = run / "revision.diff"
    lines = ["# PaperAudit revision diff", ""]
    for entry in plan.get("diff", []) or []:
        lines.append(f"## {entry.get('proposal_id', '')} / {entry.get('block_id', '')}")
        lines.append(str(entry.get("unified", "")))
        lines.append("")
    _write_json(run / "revision.plan.json", plan)
    temporary = diff_path.with_name(f".{diff_path.name}.tmp")
    temporary.write_text("\n".join(lines), encoding="utf-8")
    temporary.replace(diff_path)


def _preflight_proposals(
    selected: list[dict],
    paragraphs: dict[str, Paragraph],
    override: str | None,
) -> tuple[list[tuple[dict, list[tuple[str, int, int, str]]]], list[dict[str, Any]]]:
    """Resolve edit spans and reject unsafe/conflicting proposals atomically."""

    candidates: list[tuple[dict, list[tuple[str, int, int, str]]]] = []
    failed: list[dict[str, Any]] = []
    block_owners: dict[str, set[int]] = defaultdict(set)

    for index, finding in enumerate(selected):
        proposal = _proposal(finding)
        level = str(proposal.get("level") or finding.get("edit_level") or "").strip().casefold()
        if level in {"c", "c_substantive", "substantive", "editlevel.c_substantive"}:
            failed.append({"id": finding.get("id"), "reason": "C_substantive 修改只允许提供建议，不能自动写入"})
            continue
        targets = list(proposal.get("target_blocks") or finding.get("block_ids") or [])
        if not targets:
            failed.append({"id": finding.get("id"), "reason": "无目标位置（block_ids 为空）"})
            continue
        replacements = _replacements_for(finding, proposal, targets, override)
        if replacements is None:
            failed.append({"id": finding.get("id"), "reason": "一个意见包含多个段落时，必须为每个段落提供独立 replacement"})
            continue
        expected = _expected_texts(finding, proposal, targets)
        ranges = _ranges_for(finding, targets)
        operations: list[tuple[str, int, int, str]] = []
        error = None
        if len(set(targets)) != len(targets):
            error = "同一意见重复指定了目标段落"
        for bid in targets:
            if error:
                break
            p = paragraphs.get(bid)
            if p is None:
                error = f"定位失败：{bid}"
                break
            actual = _paragraph_text(p)
            start_end = ranges.get(bid)
            if start_end is not None:
                start, end = start_end
                if start < 0 or end < start or end > len(actual):
                    error = f"{bid}: char_range 超出段落长度"
                    break
                expected_old = expected.get(bid)
                if expected_old is None:
                    expected_old = str(finding.get("verbatim_quote") or "") or None
                if expected_old is None or actual[start:end] != expected_old:
                    error = f"{bid}: expected_old_text 与 char_range 不匹配"
                    break
            else:
                expected_old = expected.get(bid)
                if expected_old:
                    if expected_old == actual:
                        start, end = 0, len(actual)
                    elif actual.count(expected_old) == 1:
                        start = actual.index(expected_old)
                        end = start + len(expected_old)
                    else:
                        error = f"{bid}: expected_old_text 不匹配或在段落中不唯一"
                        break
                else:
                    start, end = 0, len(actual)
            operations.append((bid, start, end, replacements[bid]))
        if error:
            failed.append({"id": finding.get("id"), "reason": error})
            continue
        candidates.append((finding, operations))
        for bid in targets:
            block_owners[bid].add(len(candidates) - 1)

    conflicts: set[int] = set()
    for bid, owners in block_owners.items():
        if len(owners) > 1:
            conflicts.update(owners)
            ids = [str(candidates[owner][0].get("id")) for owner in sorted(owners)]
            failed.append({"id": ids, "reason": f"目标段落冲突：{bid}；请合并为单个修订建议"})
    return [candidate for i, candidate in enumerate(candidates) if i not in conflicts], failed


def _progress(callback: Callable[..., None] | None, stage: str, message: str, percent: int, **payload) -> None:
    if callback is not None:
        callback(stage, message, percent, **payload)


def _replace_keep_style(p: Paragraph, new_text: str) -> None:
    """Replace all visible paragraph text while retaining the paragraph XML.

    The old implementation removed every child except ``pPr``.  That also
    removed hyperlinks, OMML equations, field nodes, bookmarks and comment
    anchors.  Text nodes are edited in place instead; their containing runs
    and relationships survive the patch.
    """
    _replace_span_keep_xml(p, 0, len(_paragraph_text(p)), new_text)


def _text_nodes(p: Paragraph) -> list[Any]:
    """Return ordinary Word text nodes in document order.

    OMML ``m:t`` nodes are intentionally excluded: equations are structured
    content and must survive a prose edit.  Field instructions are also
    excluded, while field result ``w:t`` nodes remain editable.
    """
    return [
        el
        for el in p._p.iter()
        if el.tag in {_W_T, _W_DEL_TEXT}  # keep hyperlinks/bookmarks/fields
    ]


def _paragraph_text(p: Paragraph) -> str:
    return "".join(el.text or "" for el in _text_nodes(p))


def _replace_span_keep_xml(p: Paragraph, start: int, end: int, new_text: str) -> None:
    nodes = _text_nodes(p)
    old = "".join(el.text or "" for el in nodes)
    if start < 0 or end < start or end > len(old):
        raise ValueError("字符范围超出段落长度")
    if not nodes:
        if new_text:
            p.add_run(new_text)
        return

    # Edit only text nodes touched by the span.  The previous implementation
    # put the whole paragraph into the first run, which left hyperlinks and
    # field results empty even though their XML nodes survived.  Keeping the
    # prefix/suffix in their original nodes preserves unaffected formatting and
    # relationships while still allowing a cross-run replacement.
    cursor = 0
    inserted = False
    for node in nodes:
        value = node.text or ""
        node_start, node_end = cursor, cursor + len(value)
        if start == end:
            # A zero-width range is an insertion, not a replacement.  Place
            # it in the run that owns the character offset so a range inside
            # a styled run does not unexpectedly append at paragraph end.
            if not inserted and node_start <= start <= node_end:
                local = start - node_start
                node.text = value[:local] + new_text + value[local:]
                inserted = True
            cursor = node_end
            continue
        overlap_start = max(start, node_start)
        overlap_end = min(end, node_end)
        if overlap_start < overlap_end or (start == end and node_start == start):
            local_start = max(0, overlap_start - node_start)
            local_end = max(local_start, overlap_end - node_start)
            prefix = value[:local_start]
            suffix = value[local_end:]
            if not inserted:
                node.text = prefix + new_text + suffix
                inserted = True
            else:
                node.text = suffix
        cursor = node_end
    if not inserted:
        # A zero-length insertion at the end of the paragraph has no overlap.
        nodes[-1].text = (nodes[-1].text or "") + new_text


def _normalize_text(value: object) -> str:
    return " ".join(str(value or "").split())


def _proposal(finding: dict) -> dict:
    for key in ("edit_proposal", "proposal"):
        value = finding.get(key)
        if isinstance(value, dict):
            return value
    return finding


def _proposal_text(finding: dict) -> str:
    proposal = _proposal(finding)
    value = proposal.get("new_text") or proposal.get("replacement") or finding.get("suggested_fix")
    return str(value or "").strip()


def _replacements_for(
    finding: dict, proposal: dict, targets: list[str], override: str | None
) -> dict[str, str] | None:
    if override is not None and override.strip():
        if len(targets) > 1:
            return None
        return {targets[0]: override.strip()}
    value = proposal.get("replacements") or proposal.get("replacement")
    if isinstance(value, dict):
        out = {str(k): str(v) for k, v in value.items() if str(v).strip()}
        return out if all(bid in out for bid in targets) else None
    if isinstance(value, list):
        if len(value) != len(targets):
            return None
        return {bid: str(val) for bid, val in zip(targets, value)}
    value = proposal.get("new_text") or finding.get("suggested_fix")
    if not str(value or "").strip() and isinstance(proposal.get("replacement"), str):
        value = proposal.get("replacement")
    if not str(value or "").strip() or len(targets) != 1:
        return None
    return {targets[0]: str(value).strip()}


def _expected_texts(finding: dict, proposal: dict, targets: list[str]) -> dict[str, str | None]:
    value = proposal.get("expected_old_text", finding.get("expected_old_text"))
    if value is None:
        value = proposal.get("old_text", finding.get("old_text"))
    if isinstance(value, dict):
        return {bid: (str(value[bid]) if bid in value else None) for bid in targets}
    if isinstance(value, list):
        return {bid: (str(v) if v is not None else None) for bid, v in zip(targets, value)}
    return {bid: (str(value) if value is not None else None) for bid in targets}


def _ranges_for(finding: dict, targets: list[str]) -> dict[str, tuple[int, int]]:
    value = _proposal(finding).get("char_ranges", finding.get("char_ranges"))
    def _pair(v: object) -> tuple[int, int] | None:
        if not isinstance(v, (list, tuple)) or len(v) != 2:
            return None
        try:
            return int(v[0]), int(v[1])
        except (TypeError, ValueError):
            return None

    if isinstance(value, dict):
        return {str(k): pair for k, v in value.items() if (pair := _pair(v)) is not None}
    if isinstance(value, list) and value and all(isinstance(v, (list, tuple)) and len(v) == 2 for v in value):
        return {bid: pair for bid, v in zip(targets, value) if (pair := _pair(v)) is not None}
    return {}


def _load_findings(run_dir: Path) -> list[dict]:
    """读取条目。优先级：verify 产物 > 宿主原始回填 > 确定性预检。"""
    # verify 产物：{"confirmed": [...], "rejected": [...]}
    p = run_dir / "findings.json"
    if p.exists():
        data = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            if "confirmed" in data:
                return list(data["confirmed"])
            return data.get("findings", [])
        return data

    # 宿主原始回填
    p = run_dir / "findings.llm.json"
    if p.exists():
        data = json.loads(p.read_text(encoding="utf-8"))
        return data.get("findings", []) if isinstance(data, dict) else data

    p = run_dir / "manifest.json"
    if p.exists():
        data = json.loads(p.read_text(encoding="utf-8"))
        return data.get("deterministic_findings", [])
    return []


def finding_fingerprint(finding: object) -> str:
    """Return a stable instance key for regression comparison.

    IDs are deliberately omitted: verify re-numbers findings after every run.
    The detector, canonical issue type, target blocks, grounded evidence and a
    normalized subject together distinguish two instances of the same issue.
    """
    row = finding.to_dict() if hasattr(finding, "to_dict") else dict(finding or {})
    issue = str(row.get("issue_type") or "other").strip().lower()
    detector = _detector_name(row, issue)
    blocks = ",".join(sorted(str(x) for x in (row.get("block_ids") or [])))
    evidence = _normalize_text(row.get("evidence_key") or row.get("verbatim_quote"))
    subject = _normalize_text(row.get("subject") or row.get("rationale"))[:240]
    return "|".join((detector, issue, blocks, evidence, subject))


def _finding_family(finding: object) -> str:
    row = finding.to_dict() if hasattr(finding, "to_dict") else dict(finding or {})
    issue = str(row.get("issue_type") or "other").strip().lower()
    detector = _detector_name(row, issue)
    subject = _normalize_text(row.get("subject") or row.get("rationale"))[:240]
    return "|".join((detector, issue, subject))


def _detector_name(row: dict, issue: str) -> str:
    explicit = str(row.get("detector") or "").strip().lower()
    if explicit:
        return explicit
    prefixes = {
        "citation_": "citations",
        "crossref_": "crossref",
        "numeric_": "numbers",
        "stats_": "numbers",
        "terminology_": "terminology",
        "structure_": "structure",
        "privacy_": "privacy",
    }
    for prefix, detector in prefixes.items():
        if issue.startswith(prefix):
            return detector
    return str(row.get("source") or row.get("checklist_id") or "other").strip().lower()


def _finding_active(finding: object) -> bool:
    row = finding.to_dict() if hasattr(finding, "to_dict") else dict(finding or {})
    return bool(row.get("gate_passed", True)) and str(row.get("verdict", "confirmed")) != "unverifiable"


def _findings_keys(run_dir: Path) -> set[str]:
    """Compatibility helper returning fingerprint keys instead of issue types."""
    return {finding_fingerprint(f) for f in _load_findings(run_dir) if _finding_active(f)}


def _compare_findings(before: Iterable[object], after: Iterable[object]) -> dict[str, Any]:
    """Compare finding instances with multiplicity preserved.

    ``regressions`` remains a flat list for existing callers.  The structured
    report additionally exposes resolved/unchanged/new/worsened/moved counts,
    making a second instance of an existing issue visible.
    """
    before_rows = [f for f in before if _finding_active(f)]
    after_rows = [f for f in after if _finding_active(f)]
    before_counts = Counter(finding_fingerprint(f) for f in before_rows)
    after_counts = Counter(finding_fingerprint(f) for f in after_rows)
    sev_rank = {"major": 0, "minor": 1, "nit": 2}
    before_severity = {
        finding_fingerprint(f): sev_rank.get(str((f.to_dict() if hasattr(f, "to_dict") else f).get("severity", "minor")), 1)
        for f in before_rows
    }
    after_severity = {
        finding_fingerprint(f): sev_rank.get(str((f.to_dict() if hasattr(f, "to_dict") else f).get("severity", "minor")), 1)
        for f in after_rows
    }
    before_family: dict[str, list[str]] = defaultdict(list)
    after_family: dict[str, list[str]] = defaultdict(list)
    for f in before_rows:
        before_family[_finding_family(f)].append(finding_fingerprint(f))
    for f in after_rows:
        after_family[_finding_family(f)].append(finding_fingerprint(f))

    unchanged: list[dict[str, Any]] = []
    resolved: list[dict[str, Any]] = []
    worsened: list[dict[str, Any]] = []
    new: list[dict[str, Any]] = []
    moved: list[dict[str, Any]] = []
    moved_old: set[str] = set()
    moved_new: set[str] = set()

    # A changed block/evidence with the same detector/issue/subject is a move,
    # rather than simultaneously reporting one resolution and one regression.
    for family in sorted(set(before_family) & set(after_family)):
        old_keys = set(before_family[family]) - set(after_family[family])
        new_keys = set(after_family[family]) - set(before_family[family])
        for old_key, new_key in zip(sorted(old_keys), sorted(new_keys)):
            moved.append({"from": old_key, "to": new_key, "before_count": before_counts[old_key], "after_count": after_counts[new_key]})
            moved_old.add(old_key)
            moved_new.add(new_key)

    for key in sorted(set(before_counts) | set(after_counts)):
        old_n, new_n = before_counts[key], after_counts[key]
        row = {"fingerprint": key, "before_count": old_n, "after_count": new_n}
        if key in moved_old or key in moved_new:
            continue
        if old_n == new_n:
            if old_n:
                if after_severity.get(key, 1) < before_severity.get(key, 1):
                    row["before_severity"] = before_severity.get(key)
                    row["after_severity"] = after_severity.get(key)
                    worsened.append(row)
                else:
                    unchanged.append(row)
        elif new_n == 0:
            resolved.append(row)
        elif old_n == 0:
            new.append(row)
        elif new_n > old_n:
            worsened.append(row)
        else:
            resolved.append(row)

    regressions = [row["fingerprint"] for row in new + worsened]
    return {
        "resolved": resolved,
        "unchanged": unchanged,
        "new": new,
        "worsened": worsened,
        "moved": moved,
        "regressions": regressions,
        "before_counts": dict(sorted(before_counts.items())),
        "after_counts": dict(sorted(after_counts.items())),
    }
