from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from paperrevamper.evidence.citations import check as citation_check  # noqa: E402
from paperrevamper.ingest import read_docx  # noqa: E402
from paperrevamper.models import IssueType  # noqa: E402


def main() -> int:
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "author-year.docx"
        document = Document()
        document.add_paragraph("Prior work (Smith, 2024) and Smith et al. (2024) support the method.")
        document.add_paragraph("张三（2023）提供了中文背景。")
        document.add_heading("References", level=1)
        document.add_paragraph("Smith, J. (2024). A title.")
        document.add_paragraph("张三（2023）. 中文标题。")
        document.save(path)
        ir = read_docx(path)

    keys = {entry.key for entry in ir.citations}
    marks = {mark.key for mark in ir.citation_marks}
    assert "smith|2024" in keys and "smith|2024" in marks
    assert "张三|2023" in keys and "张三|2023" in marks
    findings = citation_check(ir)
    assert not any(f.issue_type in {IssueType.CITATION_MISSING, IssueType.CITATION_UNUSED} for f in findings)
    print({"status": "ok", "entries": len(ir.citations), "marks": len(ir.citation_marks)})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
