"""注入式验证：给文档埋入已知缺陷，检查确定性检测器能否召回。

这是 PLAN_v2 §七 的验收方法（seeded fixture）——比空跑真实文档更有说服力，
因为空跑返回 0 项既可能是文档规范，也可能是检测器失效。
"""

from __future__ import annotations

import os
import re
import shutil
import sys
from pathlib import Path

import docx

from paperrevamper.evidence import run_all
from paperrevamper.ingest import read_docx
from paperrevamper.reporting import apply_gate

# 源文档：优先命令行参数，其次环境变量 PAPERREVAMPER_SAMPLE，最后默认路径。
# 原开题报告样本已不在，请传入任意 .docx 作为基准。
DEFAULT_SRC = Path(r"C:\Users\mmrgr\Desktop\开题\开题报告3.3.docx")


def _resolve_src() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    env = os.environ.get("PAPERREVAMPER_SAMPLE") or os.environ.get("PAPERAUDIT_SAMPLE")
    if env:
        return Path(env)
    return DEFAULT_SRC


def build_seeded(src: Path, dst: Path) -> dict[str, bool]:
    """在副本上注入 3 类已知缺陷，返回注入成功标记。"""
    # 只复制内容，避免继承用户源文档的只读属性。
    shutil.copyfile(src, dst)
    d = docx.Document(str(dst))
    paras = [p for p in d.paragraphs]
    injected = {"missing_ref": False, "broken_crossref": False, "gap_numbering": False}

    # 1) 删掉参考文献表中的 [5] → 应触发「引而未录」
    for p in paras:
        if re.match(r"^\s*\[5\]\s*\S", p.text):
            p._p.getparent().remove(p._p)
            injected["missing_ref"] = True
            break

    # 2) 在正文中引用一个不存在的图 → 应触发交叉引用断链
    for p in paras:
        t = p.text
        if len(t) > 60 and not t.strip().startswith("[") and "参考文献" not in t:
            if p.style and "Head" in (p.style.name or ""):
                continue
            p.add_run("相关示意如图9-9所示。")
            injected["broken_crossref"] = True
            break

    # 3) 制造章节编号跳号：把某个 "2.2" 改成 "2.4"
    for p in paras:
        if p.style and "Head" in (p.style.name or "") and p.text.strip().startswith("2.2 "):
            for r in p.runs:
                if "2.2" in r.text:
                    r.text = r.text.replace("2.2", "2.4", 1)
                    injected["gap_numbering"] = True
                    break
        if injected["gap_numbering"]:
            break

    d.save(str(dst))
    return injected


def main() -> int:
    src = _resolve_src()
    if not src.exists():
        print(f"[错误] 源文档不存在：{src}", file=sys.stderr)
        print("用法：python tests/seeded_check.py <任意.docx>", file=sys.stderr)
        print("      或设置环境变量 PAPERREVAMPER_SAMPLE=<路径>", file=sys.stderr)
        return 2

    tmp = Path(__file__).parent / "_tmp"
    tmp.mkdir(exist_ok=True)
    dst = tmp / f"seeded_check_out_{os.getpid()}.docx"
    injected = build_seeded(src, dst)
    print("源文档：", src)
    print("注入结果：", injected)

    doc = read_docx(dst)
    findings = apply_gate(doc, run_all(doc))
    kinds = {f.issue_type.value for f in findings if f.gate_passed}
    print("\n检出问题类型：", sorted(kinds))
    for f in findings:
        if f.gate_passed:
            print(f"  - [{f.issue_type.value}] {f.rationale[:110]}")

    expect = {
        "missing_ref": "citation_missing",
        "broken_crossref": "crossref_broken",
        "gap_numbering": "structure_issue",
    }
    ok = True
    print("\n=== 召回验证 ===")
    for k, want in expect.items():
        hit = want in kinds
        if injected.get(k):
            print(f"  {k:16s} → 期望 {want:20s} {'命中' if hit else '未命中'}")
            if not hit:
                ok = False
        else:
            print(f"  {k:16s} → 注入失败，跳过")
    print("\n结论：", "全部召回" if ok else "存在漏检")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
