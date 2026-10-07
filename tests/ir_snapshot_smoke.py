"""Verify that prepared runs remain auditable without the original DOCX."""

from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from paperrevamper.prepare import prepare
from paperrevamper.verify import verify


def main() -> None:
    with TemporaryDirectory(prefix="paperrevamper-snapshot-") as raw:
        root = Path(raw)
        source = root / "paper.docx"
        run = root / "run"
        document = Document()
        document.add_heading("摘要", level=1)
        document.add_paragraph("结果见图1。")
        document.add_heading("参考文献", level=1)
        document.add_paragraph("[1] Smith. A paper.")
        document.save(source)

        prepare(source, run)
        source.unlink()
        result = verify(run)
        assert result["status"] == "ok"
        assert result["source_state"] == "snapshot-source-missing"
        assert (run / "ir_snapshot.json").exists()
        assert (run / "review.md").exists()
        print({"status": "ok", "source_state": result["source_state"]})


if __name__ == "__main__":
    main()
