"""Smoke checks for config-driven venue profiles and section gates."""

from __future__ import annotations

import sys
import tempfile
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document  # noqa: E402

from paperrevamper.ingest import read_docx  # noqa: E402
from paperrevamper.prepare import prepare  # noqa: E402
from paperrevamper.venue import audit_required_sections, list_venues, resolve_venue  # noqa: E402


def main() -> None:
    profiles = list_venues()
    assert profiles and resolve_venue("generic") == profiles[0]
    profile = profiles[0]
    profile_data = profile.to_dict()
    assert profile_data["stop_criteria_status"] == "advisory"
    assert profile_data["stop_criteria_evaluated"] is False
    with tempfile.TemporaryDirectory(prefix="paperrevamper-venue-") as temp:
        root = Path(temp)
        source = root / "paper.docx"
        document = Document()
        document.add_heading("摘要", level=1)
        document.add_paragraph("A short manuscript.")
        document.add_heading("引言", level=1)
        document.add_paragraph("Background.")
        document.save(source)
        ir = read_docx(source)
        missing = audit_required_sections(ir, profile)
        assert missing and all(item.gate_passed for item in missing)
        run = prepare(source, root / "run", venue="generic")
        manifest = (run / "manifest.json").read_text(encoding="utf-8")
        assert '"venue"' in manifest
        assert (run / "venue.profile.json").exists()
        manifest_data = json.loads(manifest)
        assert manifest_data["venue"]["stop_criteria_status"] == "advisory"
        assert manifest_data["venue"]["stop_criteria_evaluated"] is False
        assert any(item["checklist_id"].startswith("venue:generic-pre-submission") for item in manifest_data["deterministic_findings"])
        custom = root / "venue.yaml"
        custom.write_text("id: lab-profile\nname: Lab profile\nrequired_sections: [摘要]\n", encoding="utf-8")
        assert resolve_venue(custom) is not None
    print({"venues": len(profiles), "status": "ok"})


if __name__ == "__main__":
    main()
