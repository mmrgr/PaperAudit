from __future__ import annotations

import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from paperrevamper.ingest import read_docx  # noqa: E402


def main() -> int:
    with tempfile.TemporaryDirectory() as directory:
        original = Path(directory) / "original.docx"
        patched = Path(directory) / "notes.docx"
        document = Document()
        document.add_paragraph("Body text.")
        document.add_heading("References", level=1)
        document.add_paragraph("[1] Example reference.")
        document.save(original)
        notes_xml = b'''<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="-1"/><w:footnote w:id="1"><w:p><w:r><w:t>Footnote claim [1].</w:t></w:r></w:p></w:footnote></w:footnotes>'''
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(patched, "w") as target:
            for info in source.infolist():
                target.writestr(info, source.read(info.filename))
            target.writestr("word/footnotes.xml", notes_xml)
        ir = read_docx(patched)

    assert ir.metadata["notes"]["footnote"] == 1
    assert any(block.id.startswith("fo_") and "Footnote claim" in block.text for block in ir.blocks)
    assert any(mark.block_id.startswith("fo_") and mark.key == "1" for mark in ir.citation_marks)
    print({"status": "ok", "footnotes": ir.metadata["notes"]["footnote"]})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
