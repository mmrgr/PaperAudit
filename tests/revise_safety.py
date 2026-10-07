"""Focused regression checks for safe DOCX revision behavior."""

from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from paperrevamper.revise import _compare_findings, _replace_keep_style, apply_revision, build_revision_plan


class RevisionSafetyTests(unittest.TestCase):
    def _run(self, findings: list[dict], paragraphs: list[str] = ["old one", "old two"]):
        root = Path(tempfile.mkdtemp())
        source = root / "source.docx"
        document = Document()
        for text in paragraphs:
            document.add_paragraph(text)
        document.save(source)
        run = root / "run"
        run.mkdir()
        (run / "findings.json").write_text(json.dumps({"confirmed": findings}), encoding="utf-8")
        output = root / "revised.docx"
        return root, source, run, output

    def test_expected_old_text_is_an_optimistic_lock(self):
        _, source, run, output = self._run([{
            "id": "F1", "block_ids": ["p_0001"], "expected_old_text": "stale", "suggested_fix": "new",
        }])
        result = apply_revision(source, run, ["F1"], out_path=output, backup=False)
        self.assertEqual(result["status"], "error")
        self.assertIn("expected_old_text", result["failed"][0]["reason"])
        self.assertIsNone(result["output"])
        self.assertFalse(output.exists())
        self.assertTrue((run / "revision.plan.json").exists())
        self.assertTrue((run / "revision.diff").exists())

    def test_multi_block_fixes_require_per_block_replacements(self):
        _, source, run, output = self._run([{
            "id": "F1", "block_ids": ["p_0001", "p_0002"], "suggested_fix": "same",
        }])
        result = apply_revision(source, run, ["F1"], out_path=output, backup=False)
        self.assertEqual(result["status"], "error")
        self.assertFalse(result["applied"])

    def test_per_block_replacements_are_independent(self):
        _, source, run, output = self._run([{
            "id": "F1", "block_ids": ["p_0001", "p_0002"],
            "expected_old_text": {"p_0001": "old one", "p_0002": "old two"},
            "replacements": {"p_0001": "new one", "p_0002": "new two"},
        }])
        result = apply_revision(source, run, ["F1"], out_path=output, backup=False)
        self.assertEqual(result["status"], "ok")
        self.assertTrue(Path(result["diff"]).exists())
        self.assertEqual([p.text for p in Document(output).paragraphs], ["new one", "new two"])

    def test_structured_children_survive(self):
        document = Document()
        paragraph = document.add_paragraph("plain")
        run = paragraph.add_run("link")
        paragraph._p.remove(run._element)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), "rId9")
        hyperlink.append(run._element)
        paragraph._p.append(hyperlink)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        field_run = OxmlElement("w:r")
        field_text = OxmlElement("w:t")
        field_text.text = "1"
        field_run.append(field_text)
        field.append(field_run)
        paragraph._p.append(field)
        equation = OxmlElement("m:oMath")
        equation_run = OxmlElement("m:r")
        equation_text = OxmlElement("m:t")
        equation_text.text = "x"
        equation_run.append(equation_text)
        equation.append(equation_run)
        paragraph._p.append(equation)
        _replace_keep_style(paragraph, "replacement")
        xml = paragraph._p.xml
        self.assertIn("w:hyperlink", xml)
        self.assertIn("w:fldSimple", xml)
        self.assertIn("oMath", xml)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "structured.docx"
            document.save(path)
            with zipfile.ZipFile(path) as package:
                persisted = package.read("word/document.xml").decode("utf-8")
            self.assertIn("w:hyperlink", persisted)
            self.assertIn("w:fldSimple", persisted)
            self.assertIn("oMath", persisted)

    def test_fingerprint_comparison_preserves_counts(self):
        finding = {"issue_type": "citation_missing", "block_ids": ["p_1"], "verbatim_quote": "[3]", "rationale": "missing", "gate_passed": True}
        report = _compare_findings([finding], [finding, finding])
        self.assertEqual(report["unchanged"], [])
        self.assertEqual(len(report["worsened"]), 1)
        self.assertEqual(report["regressions"], [report["worsened"][0]["fingerprint"]])

    def test_revision_plan_is_explicit_and_non_mutating(self):
        _, source, run, output = self._run([{
            "id": "F1", "block_ids": ["p_0001"], "expected_old_text": "old one", "suggested_fix": "new one",
        }])
        before = source.read_bytes()
        plan = build_revision_plan(source, run, ["F1"])
        self.assertEqual(plan["status"], "ready")
        self.assertEqual(plan["proposals"][0]["operations"][0]["expected_old_text"], "old one")
        self.assertEqual(plan["proposals"][0]["operations"][0]["replacement"], "new one")
        self.assertTrue((run / "revision.plan.json").exists())
        self.assertTrue((run / "revision.diff").exists())
        self.assertIn("-old one", (run / "revision.diff").read_text(encoding="utf-8"))
        self.assertIn("+new one", (run / "revision.diff").read_text(encoding="utf-8"))
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse(output.exists())

    def test_revision_plan_surfaces_conflict_graph(self):
        _, source, run, _ = self._run([
            {"id": "F1", "block_ids": ["p_0001"], "expected_old_text": "old one", "suggested_fix": "new one"},
            {"id": "F2", "block_ids": ["p_0001"], "expected_old_text": "old one", "suggested_fix": "other"},
        ])
        plan = build_revision_plan(source, run, ["F1", "F2"])
        self.assertEqual(plan["status"], "blocked")
        self.assertTrue(any(edge["type"] == "target_block_conflict" for edge in plan["conflict_graph"]["edges"]))


if __name__ == "__main__":
    unittest.main()
