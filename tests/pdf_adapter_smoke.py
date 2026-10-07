"""Run a bounded PDF -> text DOCX smoke test.

PaperRevamper is intentionally DOCX-only. This script is a test adapter, not a
PDF ingestion implementation: it preserves page boundaries as plain text,
then runs the existing deterministic checks on the projected DOCX. The report
records what the projection cannot preserve (layout, tables, figures, and
author-year citation semantics).
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document

from paperrevamper.evidence import run_all
from paperrevamper.ingest import read_docx
from paperrevamper.reporting import apply_gate, to_json, to_markdown


def _extract(pdf_path: Path) -> tuple[dict, list[str]]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise SystemExit(
            "缺少 PDF 测试依赖 pdfplumber；请在运行适配测试的环境中安装它。"
        ) from exc

    pages: list[str] = []
    with pdfplumber.open(pdf_path) as pdf:
        metadata = {
            "pages": len(pdf.pages),
            "nonempty_pages": 0,
            "total_chars": 0,
            "page_chars": [],
        }
        for page in pdf.pages:
            text = page.extract_text() or ""
            pages.append(text)
            metadata["page_chars"].append(len(text))
            metadata["total_chars"] += len(text)
            if text.strip():
                metadata["nonempty_pages"] += 1
    return metadata, pages


def _project_docx(pages: list[str], output: Path) -> None:
    document = Document()
    for page_no, text in enumerate(pages, 1):
        document.add_paragraph(f"[PDF page {page_no}]")
        for line in text.splitlines():
            line = line.strip()
            if line:
                document.add_paragraph(line)
    document.save(output)


def run(pdf_path: Path, output_dir: Path) -> dict:
    if not pdf_path.exists():
        raise FileNotFoundError(pdf_path)
    output_dir.mkdir(parents=True, exist_ok=True)

    source_hash = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    metadata, pages = _extract(pdf_path)
    metadata.update(
        {
            "source": str(pdf_path),
            "sha256": source_hash,
            "adapter": "pdfplumber text extraction -> plain DOCX paragraphs",
        }
    )
    (output_dir / "pdf_metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (output_dir / "extracted.txt").write_text(
        "\n\n".join(f"=== PDF page {i} ===\n{text}" for i, text in enumerate(pages, 1)),
        encoding="utf-8",
    )

    projected = output_dir / f"{pdf_path.stem}.text-projection.docx"
    _project_docx(pages, projected)
    ir = read_docx(projected)
    findings = apply_gate(ir, run_all(ir))
    (output_dir / "findings.json").write_text(
        to_json(ir, findings), encoding="utf-8"
    )
    (output_dir / "review.md").write_text(
        to_markdown(ir, findings), encoding="utf-8"
    )

    result = {
        "status": "ok",
        "source": str(pdf_path),
        "projection": str(projected),
        "metadata": metadata,
        "docx_blocks": len(ir.blocks),
        "confirmed_findings": sum(1 for finding in findings if finding.gate_passed),
        "unverifiable_findings": sum(1 for finding in findings if not finding.gate_passed),
        "findings": str(output_dir / "findings.json"),
        "review": str(output_dir / "review.md"),
        "limitations": [
            "仅测试 PDF 文本抽取后的 DOCX 适配层，不代表 PaperRevamper 原生支持 PDF。",
            "双栏顺序、表格、图形、公式、页眉页脚和版式信息未保留。",
            "当前引用检测器主要识别 [n] 编号制，作者-年份引用可能漏检。",
            "未执行 apply，避免对从 PDF 投影出的非原生文档做改写。",
        ],
    }
    (output_dir / "result.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="PaperRevamper PDF 适配层 smoke test")
    parser.add_argument("pdf", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    result = run(args.pdf, args.out)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
