"""Smoke test for the local control panel API."""

from __future__ import annotations

import json
import os
import shutil
import socket
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path


def _ensure_fixture(path: Path) -> None:
    """Create the tiny source document when ignored local fixtures are absent.

    The smoke test is intentionally self-contained so a clean checkout does
    not depend on a developer's ignored ``tests/_tmp`` files.
    """

    if path.exists():
        return
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("研究背景", level=1)
    document.add_paragraph("结果见图1。")
    document.add_heading("研究方法", level=1)
    document.add_paragraph("采用观察性研究方法。")
    document.add_heading("参考文献", level=1)
    document.add_paragraph("[1] Example reference.")
    document.save(path)


def request(url: str, method: str = "GET", payload: dict | None = None) -> dict:
    data = json.dumps(payload).encode("utf-8") if payload is not None else None
    req = urllib.request.Request(url, data=data, method=method, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=8) as response:
        return json.loads(response.read().decode("utf-8"))


def main() -> int:
    project = Path(__file__).resolve().parents[1]
    py = sys.executable
    run_root = project / "tests" / "_tmp" / f"panel-smoke-{int(time.time())}"
    run_dir = run_root / "seeded-run"
    sample = run_root / "source.docx"
    sample.parent.mkdir(parents=True, exist_ok=True)
    fixture = project / "tests" / "_tmp" / "seeded.docx"
    _ensure_fixture(fixture)
    shutil.copy2(fixture, sample)
    original_bytes = sample.read_bytes()
    env = os.environ.copy()
    env["PYTHONPATH"] = str(project / "src")
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    base = f"http://127.0.0.1:{port}"
    proc = subprocess.Popen(
        [py, "-m", "paperaudit.cli", "panel", "--run-root", str(run_root), "--port", str(port)],
        cwd=project,
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    try:
        for _ in range(30):
            try:
                request(base + "/api/health")
                break
            except (urllib.error.URLError, ConnectionError):
                time.sleep(0.1)
        else:
            raise RuntimeError("panel server did not start")

        health = request(base + "/api/health")
        request(base + "/api/workflow/catalog")
        custom = request(base + "/api/workflow")["workflow"]
        next(role for role in custom["roles"] if role["id"] == "language_editor")["enabled"] = False
        custom["roles"].append({
            "id": "methods_gate",
            "label": "自定义方法门禁",
            "enabled": True,
            "order": 2,
            "checklist_groups": ["method"],
            "skills": ["paper_audit"],
            "prompt": "检查方法门禁。",
            "position": {"x": 460, "y": 90},
        })
        custom["task_overrides"] = {"review:method_data_reviewer": {"depends_on": ["review:argument_reviewer"]}}
        custom["feedback_edges"] = [{"source": "verify", "target": "review:argument_reviewer", "max_iterations": 2}]
        custom["canvas"] = {"start": {"label": "自定义入口"}, "end": {"label": "自定义报告"}}
        validated = request(base + "/api/workflow/validate", "POST", {"workflow": custom})
        assert validated["status"] == "ok"
        methods_role = next(role for role in validated["workflow"]["roles"] if role["id"] == "methods_gate")
        assert methods_role["position"] == {"x": 460, "y": 90}
        assert validated["workflow"]["feedback_edges"] == [{"source": "verify", "target": "review:argument_reviewer", "max_iterations": 2}]
        assert validated["workflow"]["canvas"]["start"]["label"] == "自定义入口"
        plan = __import__("paperaudit.collaboration", fromlist=["build_plan"]).build_plan("paper", {"groups": [{"id": "method"}, {"id": "argument"}]}, validated["workflow"])
        assert next(task for task in plan["tasks"] if task["id"] == "review:method_data_reviewer")["depends_on"] == ["review:argument_reviewer"]
        methods_task = next(task for task in plan["tasks"] if task["id"] == "review:methods_gate")
        assert methods_task["depends_on"] == ["review:argument_reviewer", "review:method_data_reviewer"]
        assert plan["feedback_edges"][0]["target"] == "review:argument_reviewer"
        invalid = json.loads(json.dumps(validated["workflow"]))
        invalid["task_overrides"] = {"verify": {"depends_on": ["revise"]}, "revise": {"depends_on": ["verify"]}}
        try:
            request(base + "/api/workflow/validate", "POST", {"workflow": invalid})
        except urllib.error.HTTPError as exc:
            assert exc.code == 400
        else:
            raise AssertionError("cyclic workflow was accepted")
        deleted = json.loads(json.dumps(validated["workflow"]))
        deleted["roles"] = [role for role in deleted["roles"] if role["id"] != "methods_gate"]
        deleted_ok = request(base + "/api/workflow/validate", "POST", {"workflow": deleted})
        assert not any(role["id"] == "methods_gate" for role in deleted_ok["workflow"]["roles"])
        saved = request(base + "/api/workflow/save", "POST", {"name": "custom-smoke.json", "workflow": validated["workflow"]})
        assert saved["status"] == "ok" and Path(saved["path"]).exists()
        prepared = request(
            base + "/api/prepare",
            "POST",
            {"source": str(sample), "out_dir": str(run_dir), "workflow": validated["workflow"]},
        )
        try:
            request(base + "/api/review/start", "POST", {"run_dir": str(run_dir), "profile_id": "host_agent"})
        except urllib.error.HTTPError as exc:
            assert exc.code == 400
        else:
            raise AssertionError("host_agent was incorrectly accepted as direct model runner")
        for invalid_panel in (
            {"run_dir": str(run_dir), "profiles": ["openai"]},
            {"run_dir": str(run_dir), "profiles": ["host_agent", "openai"]},
        ):
            try:
                request(base + "/api/panel/start", "POST", invalid_panel)
            except urllib.error.HTTPError as exc:
                assert exc.code == 400
            else:
                raise AssertionError("invalid panel profile selection was accepted")
        try:
            request(
                base + "/api/prepare",
                "POST",
                {"source": str(sample), "out_dir": str(run_dir / "invalid-parser"), "pdf_parser": "invalid"},
            )
        except urllib.error.HTTPError as exc:
            assert exc.code == 400
        else:
            raise AssertionError("invalid PDF parser was accepted")
        state = request(base + "/api/state?run_dir=" + urllib.parse.quote(str(run_dir)))
        assert health["status"] == "ok"
        assert prepared["source_hash"] == state["source_hash"]
        assert state["roles"]
        assert next(stage for stage in state["stages"] if stage["id"] == "verify")["depends_on"] == ["review:methods_gate"]
        assert any(stage["id"] == "verify" for stage in state["stages"])
        assert state["trace_cursor"] >= 2
        manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
        deterministic = next(item for item in manifest["deterministic_findings"] if item.get("verbatim_quote"))
        reviewer = state["roles"][0]
        role_output = {
            "findings": [{
                "checklist_id": "A04",
                "severity": "minor",
                "block_ids": [deterministic["block_ids"][0]],
                "verbatim_quote": deterministic["verbatim_quote"],
                "rationale": "panel smoke finding",
                "confidence": 0.8,
                "reviewer_id": reviewer["id"],
                "suggested_fix": deterministic["verbatim_quote"] + "（面板测试）",
            }]
        }
        (run_dir / reviewer["output_file"]).write_text(json.dumps(role_output, ensure_ascii=False), encoding="utf-8")
        verified = request(base + "/api/verify", "POST", {"run_dir": str(run_dir)})
        state = verified["state"]
        candidate = next(item for item in state["findings"]["confirmed"] if item.get("reviewer_id") == reviewer["id"])
        panel_judgments = {
            "judgments": [
                {"finding_id": candidate["id"], "judge_model": model, "position": position, "verdict": "yes"}
                for model in ("panel-a", "panel-b")
                for position in ("claim_first", "evidence_first")
            ]
        }
        (run_dir / "panel.judgments.json").write_text(json.dumps(panel_judgments), encoding="utf-8")
        adjudicated = request(base + "/api/adjudicate", "POST", {"run_dir": str(run_dir), "judgments": "panel.judgments.json"})
        assert adjudicated["result"]["summary"]["confirmed"] == 1
        assert next(item for item in adjudicated["state"]["findings"]["confirmed"] if item["id"] == candidate["id"])["panel_verdict"] == "confirmed"
        try:
            request(base + "/api/apply", "POST", {"run_dir": str(run_dir), "source": str(sample), "source_hash": state["source_hash"], "finding_ids": [candidate["id"]], "confirm": True})
        except urllib.error.HTTPError as exc:
            assert exc.code == 400
        else:
            raise AssertionError("unaccepted finding was applied")
        request(base + "/api/decisions", "POST", {"run_dir": str(run_dir), "finding_id": candidate["id"], "decision": "accept", "source_hash": state["source_hash"]})
        try:
            request(base + "/api/apply/start", "POST", {"run_dir": str(run_dir), "source": str(sample), "source_hash": state["source_hash"], "finding_ids": [candidate["id"]], "confirm": True})
        except urllib.error.HTTPError as exc:
            assert exc.code == 400
        else:
            raise AssertionError("panel apply bypassed the required revision plan preview")
        planned = request(
            base + "/api/revision/plan",
            "POST",
            {"run_dir": str(run_dir), "source": str(sample), "finding_ids": [candidate["id"]]},
        )
        assert planned["result"]["status"] == "ready"
        assert Path(planned["result"]["revision_plan"]).exists()
        assert planned["result"]["proposals"][0]["operations"]
        output = run_dir / "revised.docx"
        applied = request(base + "/api/apply", "POST", {"run_dir": str(run_dir), "source": str(sample), "source_hash": state["source_hash"], "finding_ids": [candidate["id"]], "out": str(output), "confirm": True})
        assert applied["result"]["status"] == "ok"
        assert output.exists()
        assert applied["result"]["regressions"] == []
        assert Path(applied["result"]["backup"]).exists()
        assert sample.read_bytes() == original_bytes
        print(json.dumps({"status": "ok", "roles": len(state["roles"]), "stages": len(state["stages"]), "trace": applied["state"]["trace_cursor"], "applied": len(applied["result"]["applied"])}))
        return 0
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=3)
        except subprocess.TimeoutExpired:
            proc.kill()


if __name__ == "__main__":
    raise SystemExit(main())
