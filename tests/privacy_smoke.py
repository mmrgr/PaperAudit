"""Small deterministic smoke test for the DOCX privacy auditor."""

from __future__ import annotations

import json
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from paperaudit.evidence.privacy import check
from paperaudit.ingest import read_docx
from paperaudit.models import IssueType


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="paperaudit-privacy-") as raw:
        path = Path(raw) / "paper.docx"
        document = Document()
        document.core_properties.author = "Alice Example"
        document.add_paragraph("作者：Alice Example，邮箱 alice@example.com")
        document.sections[0].header.paragraphs[0].text = "Alice Example Lab"
        document.save(path)

        # Add representative OOXML-only identity leaks: a comment author,
        # tracked change author, and hidden text that are absent from the IR.
        temp = path.with_suffix(".patched.docx")
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename == "word/document.xml":
                    marker = b"</w:body>"
                    hidden = (
                        b'<w:p><w:r><w:rPr><w:vanish/></w:rPr>'
                        b'<w:t>hidden identity</w:t></w:r></w:p>'
                    )
                    payload = payload.replace(marker, hidden + marker)
                    payload = payload.replace(
                        b"</w:body>",
                        b'<w:ins w:id="9" w:author="Alice Example" w:date="2026-01-01T00:00:00Z"/></w:body>',
                    )
                target.writestr(info, payload)
            target.writestr(
                "word/comments.xml",
                b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b'<w:comment w:id="0" w:author="Alice Example"><w:p><w:r><w:t>todo</w:t></w:r></w:p></w:comment>'
                b"</w:comments>",
            )
        temp.replace(path)

        findings = check(read_docx(path))
        assert findings
        assert all(f.issue_type is IssueType.PRIVACY_RISK for f in findings)
        kinds = " ".join(f.rationale for f in findings)
        for needle in ("核心属性", "批注", "修订", "隐藏文字", "邮箱"):
            assert needle in kinds, (needle, kinds)
        print(json.dumps({"status": "ok", "findings": len(findings)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
